#!/usr/bin/env python3
"""Inspect Bento PPTX packages for completion evidence."""

from __future__ import annotations

import argparse
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation


PML = "http://schemas.openxmlformats.org/presentationml/2006/main"


def xml_shape_report(slide_xml: bytes) -> dict:
    root = ET.fromstring(slide_xml)
    shape_count = len(root.findall(f".//{{{PML}}}sp"))
    picture_count = len(root.findall(f".//{{{PML}}}pic"))
    group_count = len(root.findall(f".//{{{PML}}}grpSp"))
    graphic_frame_count = len(root.findall(f".//{{{PML}}}graphicFrame"))
    return {
        "native_shape_count": shape_count + picture_count + group_count + graphic_frame_count,
        "shape_count": shape_count,
        "picture_count": picture_count,
        "group_count": group_count,
        "graphic_frame_count": graphic_frame_count,
    }


def inspect_pptx(path: Path) -> dict:
    with ZipFile(path) as package:
        names = package.namelist()
        slide_xml = sorted(
            name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        xml_reports = [xml_shape_report(package.read(name)) for name in slide_xml]
        svg_media = sorted(name for name in names if name.startswith("ppt/media/") and name.endswith(".svg"))
        png_media = sorted(name for name in names if name.startswith("ppt/media/") and name.endswith(".png"))
        content_types = package.read("[Content_Types].xml").decode("utf-8")
    prs = Presentation(str(path))
    return {
        "path": str(path),
        "slide_count_zip": len(slide_xml),
        "slide_count_pptx": len(prs.slides),
        "shape_counts": [len(slide.shapes) for slide in prs.slides],
        "native_shape_counts": [item["native_shape_count"] for item in xml_reports],
        "picture_counts": [item["picture_count"] for item in xml_reports],
        "group_counts": [item["group_count"] for item in xml_reports],
        "svg_media_count": len(svg_media),
        "png_media_count": len(png_media),
        "has_svg_content_type": "image/svg+xml" in content_types,
        "title": prs.core_properties.title,
        "size_bytes": path.stat().st_size,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect a Bento PPTX package.")
    parser.add_argument("pptx", help="PPTX file to inspect")
    parser.add_argument("--expect-slides", type=int, help="Expected slide count")
    parser.add_argument("--require-svg-media", action="store_true", help="Fail unless SVG media exists for every slide")
    parser.add_argument("--require-png-media", action="store_true", help="Fail unless PNG media exists for every slide")
    parser.add_argument("--require-native-editable", action="store_true", help="Fail unless each slide has multiple native PPT objects")
    parser.add_argument("--min-native-shapes-per-slide", type=int, default=2, help="Minimum XML shape/picture/group objects per slide")
    args = parser.parse_args()

    path = Path(args.pptx).expanduser().resolve()
    report = inspect_pptx(path)
    errors: list[str] = []

    if args.expect_slides is not None:
        if report["slide_count_zip"] != args.expect_slides or report["slide_count_pptx"] != args.expect_slides:
            errors.append(f"Expected {args.expect_slides} slides")
    if args.require_svg_media and report["svg_media_count"] < report["slide_count_zip"]:
        errors.append("Missing SVG media for one or more slides")
    if args.require_svg_media and not report["has_svg_content_type"]:
        errors.append("Missing SVG content type")
    if args.require_png_media and report["png_media_count"] < report["slide_count_zip"]:
        errors.append("Missing PNG media for one or more slides")
    if args.require_native_editable:
        too_sparse = [
            index + 1
            for index, count in enumerate(report["native_shape_counts"])
            if count < args.min_native_shapes_per_slide
        ]
        if too_sparse:
            errors.append(
                "Slides do not contain enough native editable objects: "
                + ", ".join(str(index) for index in too_sparse)
            )
    if any(count == 0 for count in report["shape_counts"]):
        errors.append("One or more slides contain no shapes")

    report["status"] = "fail" if errors else "pass"
    report["errors"] = errors
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
