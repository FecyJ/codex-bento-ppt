#!/usr/bin/env python3
"""Assemble SVG slides into editable and snapshot PowerPoint decks."""

from __future__ import annotations

import argparse
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import cairosvg
from pptx import Presentation
from pptx.util import Emu

from ppt_master_bridge import PptMasterBridgeError, load_native_builder


CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CORE = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC = "http://purl.org/dc/elements/1.1/"

ET.register_namespace("", CONTENT_TYPES)
ET.register_namespace("p", PML)
ET.register_namespace("a", DRAWING)
ET.register_namespace("r", REL)
ET.register_namespace("cp", CORE)
ET.register_namespace("dc", DC)


def sorted_svgs(svg_dir: Path) -> list[Path]:
    files = sorted(svg_dir.glob("*.svg"))
    if not files:
        raise SystemExit(f"No SVG files found in {svg_dir}")
    return files


def make_base_pptx(svg_files: list[Path], pptx_path: Path, title: str) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    prs.core_properties.title = title
    blank = prs.slide_layouts[6]
    while len(prs.slides) < len(svg_files):
        prs.slides.add_slide(blank)
    prs.save(pptx_path)


def add_svg_content_type(root: ET.Element) -> None:
    for item in root.findall(f"{{{CONTENT_TYPES}}}Default"):
        if item.attrib.get("Extension") == "svg":
            item.attrib["ContentType"] = "image/svg+xml"
            return
    ET.SubElement(root, f"{{{CONTENT_TYPES}}}Default", {"Extension": "svg", "ContentType": "image/svg+xml"})


def next_relationship_id(root: ET.Element) -> str:
    max_id = 0
    for rel in root.findall(f"{{{PKG_RELS}}}Relationship"):
        raw = rel.attrib.get("Id", "")
        if raw.startswith("rId"):
            try:
                max_id = max(max_id, int(raw[3:]))
            except ValueError:
                pass
    return f"rId{max_id + 1}"


def append_picture(slide_root: ET.Element, rel_id: str, shape_id: int, name: str) -> None:
    sp_tree = slide_root.find(f".//{{{PML}}}spTree")
    if sp_tree is None:
        raise RuntimeError("Slide XML has no p:spTree")

    pic = ET.Element(f"{{{PML}}}pic")
    nv_pic_pr = ET.SubElement(pic, f"{{{PML}}}nvPicPr")
    ET.SubElement(nv_pic_pr, f"{{{PML}}}cNvPr", {"id": str(shape_id), "name": name})
    c_nv_pic = ET.SubElement(nv_pic_pr, f"{{{PML}}}cNvPicPr")
    ET.SubElement(c_nv_pic, f"{{{DRAWING}}}picLocks", {"noChangeAspect": "1"})
    ET.SubElement(nv_pic_pr, f"{{{PML}}}nvPr")

    blip_fill = ET.SubElement(pic, f"{{{PML}}}blipFill")
    ET.SubElement(blip_fill, f"{{{DRAWING}}}blip", {f"{{{REL}}}embed": rel_id})
    stretch = ET.SubElement(blip_fill, f"{{{DRAWING}}}stretch")
    ET.SubElement(stretch, f"{{{DRAWING}}}fillRect")

    sp_pr = ET.SubElement(pic, f"{{{PML}}}spPr")
    xfrm = ET.SubElement(sp_pr, f"{{{DRAWING}}}xfrm")
    ET.SubElement(xfrm, f"{{{DRAWING}}}off", {"x": "0", "y": "0"})
    ET.SubElement(xfrm, f"{{{DRAWING}}}ext", {"cx": "12192000", "cy": "6858000"})
    geom = ET.SubElement(sp_pr, f"{{{DRAWING}}}prstGeom", {"prst": "rect"})
    ET.SubElement(geom, f"{{{DRAWING}}}avLst")
    sp_tree.append(pic)


def embed_svgs(base_pptx: Path, svg_files: list[Path], output_pptx: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        with zipfile.ZipFile(base_pptx, "r") as zin:
            zin.extractall(tmp_dir)

        content_types_path = tmp_dir / "[Content_Types].xml"
        content_root = ET.parse(content_types_path).getroot()
        add_svg_content_type(content_root)
        ET.ElementTree(content_root).write(content_types_path, encoding="utf-8", xml_declaration=True)

        media_dir = tmp_dir / "ppt" / "media"
        media_dir.mkdir(exist_ok=True)

        for index, svg_path in enumerate(svg_files, start=1):
            media_name = f"codex_bento_slide_{index:02d}.svg"
            shutil.copyfile(svg_path, media_dir / media_name)

            rels_path = tmp_dir / "ppt" / "slides" / "_rels" / f"slide{index}.xml.rels"
            rel_tree = ET.parse(rels_path)
            rel_root = rel_tree.getroot()
            rel_id = next_relationship_id(rel_root)
            ET.SubElement(
                rel_root,
                f"{{{PKG_RELS}}}Relationship",
                {
                    "Id": rel_id,
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                    "Target": f"../media/{media_name}",
                },
            )
            rel_tree.write(rels_path, encoding="utf-8", xml_declaration=True)

            slide_path = tmp_dir / "ppt" / "slides" / f"slide{index}.xml"
            slide_tree = ET.parse(slide_path)
            append_picture(slide_tree.getroot(), rel_id, 1000 + index, f"Codex Bento SVG {index:02d}")
            slide_tree.write(slide_path, encoding="utf-8", xml_declaration=True)

        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_pptx, "w", zipfile.ZIP_DEFLATED) as zout:
            for path in tmp_dir.rglob("*"):
                if path.is_file():
                    zout.write(path, path.relative_to(tmp_dir).as_posix())


def build_png_fallback(svg_files: list[Path], output_pptx: Path, title: str) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    prs.core_properties.title = title
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        for index, svg_path in enumerate(svg_files, start=1):
            png_path = tmp_dir / f"slide_{index:02d}.png"
            cairosvg.svg2png(url=str(svg_path), write_to=str(png_path), output_width=1280, output_height=720)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_pptx)


def set_core_title(pptx_path: Path, title: str) -> None:
    """Set docProps/core.xml title without asking python-pptx to rewrite slides."""
    if not title:
        return

    tmp_path = pptx_path.with_suffix(pptx_path.suffix + ".tmp")
    with zipfile.ZipFile(pptx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        wrote_core = False
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                root = ET.fromstring(data)
                title_el = root.find(f"{{{DC}}}title")
                if title_el is None:
                    title_el = ET.SubElement(root, f"{{{DC}}}title")
                title_el.text = title
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                wrote_core = True
            zout.writestr(item, data)

        if not wrote_core:
            root = ET.Element(f"{{{CORE}}}coreProperties")
            title_el = ET.SubElement(root, f"{{{DC}}}title")
            title_el.text = title
            zout.writestr(
                "docProps/core.xml",
                ET.tostring(root, encoding="utf-8", xml_declaration=True),
            )
    tmp_path.replace(pptx_path)


def build_editable_native(
    svg_files: list[Path],
    output_pptx: Path,
    title: str,
    ppt_master_skill_dir: str | None,
    merge_paragraphs: bool,
    conversion_trace_path: Path | None,
    quiet: bool,
) -> None:
    """Create a PPTX whose SVG elements are native DrawingML objects."""
    try:
        create_pptx_with_native_svg = load_native_builder(ppt_master_skill_dir)
    except PptMasterBridgeError as exc:
        raise SystemExit(str(exc)) from exc

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    ok = create_pptx_with_native_svg(
        svg_files=svg_files,
        output_path=output_pptx,
        canvas_format=None,
        verbose=not quiet,
        transition=None,
        use_native_shapes=True,
        merge_paragraphs=merge_paragraphs,
        conversion_trace_path=conversion_trace_path,
    )
    if not ok:
        raise SystemExit("Native editable PPTX export failed")
    set_core_title(output_pptx, title)


def default_snapshot_path(output_pptx: Path) -> Path:
    return output_pptx.with_name(f"{output_pptx.stem}_snapshot{output_pptx.suffix}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Assemble SVG slides into editable PPTX.")
    parser.add_argument("svg_dir", help="Directory containing slide SVG files")
    parser.add_argument("output_pptx", help="Output editable native PPTX path")
    parser.add_argument("--title", default="Codex Bento PPT", help="Deck title metadata")
    parser.add_argument(
        "--snapshot-output",
        help="Snapshot PPTX path. Defaults to <output_stem>_snapshot.pptx unless --no-snapshot is used.",
    )
    parser.add_argument("--no-snapshot", action="store_true", help="Only create the editable native PPTX")
    parser.add_argument("--only-snapshot", action="store_true", help="Skip native export and write only the snapshot deck")
    parser.add_argument("--png-fallback", action="store_true", help="Use raster PNG pages for the snapshot deck")
    parser.add_argument("--merge-paragraphs", action="store_true", help="Merge supported paragraph lines into larger editable text boxes")
    parser.add_argument(
        "--conversion-trace",
        nargs="?",
        const="",
        help="Write native conversion diagnostics JSON. Optional value overrides the default <output>.trace.json path.",
    )
    parser.add_argument("--ppt-master-skill-dir", help="Path to the ppt-master skill directory")
    parser.add_argument("--quiet", action="store_true", help="Reduce converter output")
    args = parser.parse_args()

    svg_files = sorted_svgs(Path(args.svg_dir).expanduser().resolve())
    output_pptx = Path(args.output_pptx).expanduser().resolve()

    trace_path: Path | None = None
    if args.conversion_trace is not None:
        trace_path = (
            output_pptx.with_suffix(".trace.json")
            if args.conversion_trace == ""
            else Path(args.conversion_trace).expanduser().resolve()
        )

    if not args.only_snapshot:
        build_editable_native(
            svg_files=svg_files,
            output_pptx=output_pptx,
            title=args.title,
            ppt_master_skill_dir=args.ppt_master_skill_dir,
            merge_paragraphs=args.merge_paragraphs,
            conversion_trace_path=trace_path,
            quiet=args.quiet,
        )
        print(f"Wrote editable native deck with {len(svg_files)} slides to {output_pptx}")

    snapshot_output = None
    if not args.no_snapshot:
        snapshot_output = (
            Path(args.snapshot_output).expanduser().resolve()
            if args.snapshot_output
            else (output_pptx if args.only_snapshot else default_snapshot_path(output_pptx))
        )
        if args.png_fallback:
            build_png_fallback(svg_files, snapshot_output, args.title)
            print(f"Wrote PNG snapshot deck with {len(svg_files)} slides to {snapshot_output}")
        else:
            with tempfile.TemporaryDirectory() as tmp:
                base_pptx = Path(tmp) / "base.pptx"
                make_base_pptx(svg_files, base_pptx, args.title)
                embed_svgs(base_pptx, svg_files, snapshot_output)
            print(f"Wrote SVG snapshot deck with {len(svg_files)} slides to {snapshot_output}")

    if args.only_snapshot and snapshot_output is None:
        raise SystemExit("--only-snapshot requires snapshot output; remove --no-snapshot")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
