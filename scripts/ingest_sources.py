#!/usr/bin/env python3
"""Read user-provided raw materials for a Codex Bento PPT project."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
import shutil
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path


SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".docx", ".pdf", ".xlsx", ".xlsm", ".csv", ".tsv"}
ASSET_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg", ".bmp", ".tif", ".tiff"}


@dataclass
class SourceRecord:
    path: str
    type: str
    title: str
    status: str
    char_count: int
    headings: list[str]
    topics: list[str]
    reusable_assets: list[str]
    warnings: list[str]
    excerpt: str
    extraction_method: str


def iter_source_files(paths: list[Path]) -> list[Path]:
    files: list[Path] = []
    for raw_path in paths:
        path = raw_path.expanduser().resolve()
        if path.is_dir():
            for item in sorted(path.rglob("*")):
                if item.is_file() and not any(part.startswith(".") for part in item.parts):
                    files.append(item)
        elif path.is_file():
            files.append(path)
        else:
            files.append(path)
    return files


def read_text_file(path: Path) -> tuple[str, list[str]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    headings = [match.group(1).strip() for match in re.finditer(r"^#{1,6}\s+(.+)$", text, re.MULTILINE)]
    return text, headings


def asset_name(source_path: Path, raw_name: str) -> str:
    digest = hashlib.sha1(str(source_path).encode("utf-8")).hexdigest()[:10]
    clean = re.sub(r"[^A-Za-z0-9._-]+", "_", raw_name).strip("_") or "asset"
    return f"{source_path.stem}_{digest}_{clean}"


def copy_asset(path: Path, asset_dir: Path) -> str:
    asset_dir.mkdir(parents=True, exist_ok=True)
    output = asset_dir / asset_name(path, path.name)
    if output.resolve() != path.resolve():
        shutil.copy2(path, output)
    return str(output)


def extract_zip_media(path: Path, asset_dir: Path, media_prefix: str) -> list[str]:
    extracted: list[str] = []
    if not zipfile.is_zipfile(path):
        return extracted
    asset_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as package:
        for name in package.namelist():
            if not name.startswith(media_prefix) or name.endswith("/"):
                continue
            suffix = Path(name).suffix.lower()
            if suffix not in ASSET_EXTENSIONS:
                continue
            output = asset_dir / asset_name(path, Path(name).name)
            output.write_bytes(package.read(name))
            extracted.append(str(output))
    return extracted


def read_docx(path: Path, asset_dir: Path) -> tuple[str, list[str], list[str]]:
    from docx import Document

    document = Document(str(path))
    lines: list[str] = []
    headings: list[str] = []
    assets: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        lines.append(text)
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            headings.append(text)
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            lines.append(f"[Table {index}]\n" + "\n".join(rows))
            assets.append(f"table_{index}")
    assets.extend(extract_zip_media(path, asset_dir, "word/media/"))
    return "\n\n".join(lines), headings, assets


def read_pdf(path: Path, asset_dir: Path) -> tuple[str, list[str], list[str]]:
    import pdfplumber

    logging.getLogger("pdfminer").setLevel(logging.ERROR)
    lines: list[str] = []
    assets: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                lines.append(f"[Page {index}]\n{text.strip()}")
            tables = page.extract_tables() or []
            for table_index, table in enumerate(tables, start=1):
                assets.append(f"page_{index}_table_{table_index}")
    try:
        from pypdf import PdfReader

        asset_dir.mkdir(parents=True, exist_ok=True)
        reader = PdfReader(str(path))
        for page_index, page in enumerate(reader.pages, start=1):
            for image_index, image in enumerate(getattr(page, "images", []), start=1):
                raw_name = getattr(image, "name", "") or f"page_{page_index}_image_{image_index}.bin"
                suffix = Path(raw_name).suffix or ".bin"
                output = asset_dir / asset_name(path, f"page_{page_index}_image_{image_index}{suffix}")
                output.write_bytes(image.data)
                assets.append(str(output))
    except Exception:
        pass
    headings = [line.strip() for line in lines if len(line.strip()) < 80 and not line.startswith("[Page")][:20]
    return "\n\n".join(lines), headings, assets


def read_spreadsheet(path: Path, asset_dir: Path) -> tuple[str, list[str], list[str]]:
    import pandas as pd

    assets: list[str] = []
    ext = path.suffix.lower()
    if ext in {".csv", ".tsv"}:
        separator = "\t" if ext == ".tsv" else ","
        frame = pd.read_csv(path, sep=separator)
        preview = frame.head(40).to_markdown(index=False)
        headings = list(map(str, frame.columns[:20]))
        return f"[Sheet: {path.stem}]\n{preview}", headings, assets

    sheets = pd.read_excel(path, sheet_name=None)
    sections: list[str] = []
    headings: list[str] = []
    for sheet_name, frame in sheets.items():
        headings.append(str(sheet_name))
        preview = frame.head(40).to_markdown(index=False)
        sections.append(f"[Sheet: {sheet_name}]\n{preview}")
    assets.extend(extract_zip_media(path, asset_dir, "xl/media/"))
    return "\n\n".join(sections), headings, assets


def infer_title(path: Path, text: str, headings: list[str]) -> str:
    if headings:
        return headings[0]
    for line in text.splitlines():
        cleaned = line.strip()
        if cleaned:
            return cleaned[:120]
    return path.stem


def infer_topics(text: str, headings: list[str]) -> list[str]:
    candidates = headings[:8]
    if not candidates:
        words = re.findall(r"[\u4e00-\u9fffA-Za-z0-9][\u4e00-\u9fffA-Za-z0-9 _-]{2,30}", text)
        seen: set[str] = set()
        for word in words:
            cleaned = word.strip()
            if cleaned and cleaned not in seen:
                seen.add(cleaned)
                candidates.append(cleaned)
            if len(candidates) >= 8:
                break
    return candidates


def clean_excerpt(text: str, limit: int = 900) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    return compact[:limit]


def extract_one(path: Path, asset_dir: Path) -> SourceRecord:
    warnings: list[str] = []
    assets: list[str] = []
    ext = path.suffix.lower()
    if not path.exists():
        return SourceRecord(str(path), ext or "unknown", path.name, "missing", 0, [], [], [], ["path does not exist"], "", "fallback-script")
    if ext in ASSET_EXTENSIONS:
        copied_asset = copy_asset(path, asset_dir)
        return SourceRecord(
            str(path),
            ext.lstrip("."),
            path.name,
            "asset",
            0,
            [],
            [],
            [copied_asset],
            ["visual asset copied to source_assets; text not extracted"],
            "",
            "fallback-script",
        )
    if ext not in SUPPORTED_EXTENSIONS:
        return SourceRecord(str(path), ext or "unknown", path.name, "unsupported", 0, [], [], [], ["unsupported extension"], "", "fallback-script")

    try:
        if ext in {".md", ".markdown", ".txt"}:
            text, headings = read_text_file(path)
        elif ext == ".docx":
            text, headings, assets = read_docx(path, asset_dir)
        elif ext == ".pdf":
            text, headings, assets = read_pdf(path, asset_dir)
        elif ext in {".xlsx", ".xlsm", ".csv", ".tsv"}:
            text, headings, assets = read_spreadsheet(path, asset_dir)
        else:
            text, headings = "", []
    except Exception as exc:  # noqa: BLE001
        return SourceRecord(str(path), ext, path.name, "failed", 0, [], [], [], [f"extraction failed: {exc}"], "", "fallback-script")

    if not text.strip():
        warnings.append("no text extracted")
    title = infer_title(path, text, headings)
    topics = infer_topics(text, headings)
    return SourceRecord(
        path=str(path),
        type=ext.lstrip("."),
        title=title,
        status="extracted",
        char_count=len(text),
        headings=headings[:30],
        topics=topics,
        reusable_assets=assets,
        warnings=warnings,
        excerpt=clean_excerpt(text),
        extraction_method="fallback-script",
    )


def write_digest(records: list[SourceRecord], output_path: Path) -> None:
    extracted = [record for record in records if record.status == "extracted"]
    assets = [record for record in records if record.status == "asset"]
    other = [record for record in records if record.status not in {"extracted", "asset"}]

    lines = ["# Source Digest", "", "## Source Coverage", "", "### Extracted Text Sources", ""]
    for record in extracted:
        lines.append(f"- `{record.path}`: {record.status}, {record.type}, {record.char_count} chars, title: {record.title}")
    if not extracted:
        lines.append("- None.")

    lines.extend(["", "### Visual Or Asset Sources", ""])
    for record in assets:
        lines.append(f"- `{record.path}`: {record.type}, title: {record.title}")
    if not assets:
        lines.append("- None detected.")

    lines.extend(["", "### Unsupported Or Failed Sources", ""])
    for record in other:
        warning = "; ".join(record.warnings) if record.warnings else record.status
        lines.append(f"- `{record.path}`: {record.type}, {warning}")
    if not other:
        lines.append("- None.")

    lines.extend(["", "## Main Topics", ""])
    seen: set[str] = set()
    for record in extracted:
        for topic in record.topics:
            if topic not in seen:
                seen.add(topic)
                lines.append(f"- {topic}")
    if not seen:
        lines.append("- None extracted.")

    lines.extend(["", "## Evidence Candidates", ""])
    for record in extracted:
        if record.excerpt:
            lines.append(f"### {record.title}")
            lines.append("")
            lines.append(f"- Source: `{record.path}`")
            lines.append(f"- Excerpt: {record.excerpt}")
            lines.append("")
    lines.extend(["## Reusable Visual Material", ""])
    any_asset = False
    for record in records:
        for asset in record.reusable_assets:
            any_asset = True
            lines.append(f"- `{record.path}`: {asset}")
    if not any_asset:
        lines.append("- No reusable tables or figures detected by automatic extraction.")
    lines.extend(["", "## Extraction Warnings", ""])
    any_warning = False
    for record in records:
        for warning in record.warnings:
            any_warning = True
            lines.append(f"- `{record.path}`: {warning}")
    if not any_warning:
        lines.append("- None.")
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Ingest raw source material for a Codex Bento PPT project.")
    parser.add_argument("source_paths", nargs="+", help="Files or directories supplied by the user")
    parser.add_argument("--project-dir", required=True, help="Project directory that receives source_inventory.json and source_digest.md")
    args = parser.parse_args()

    project_dir = Path(args.project_dir).expanduser().resolve()
    project_dir.mkdir(parents=True, exist_ok=True)
    files = iter_source_files([Path(item) for item in args.source_paths])
    asset_dir = project_dir / "source_assets"
    asset_dir.mkdir(exist_ok=True)
    records = [extract_one(path, asset_dir) for path in files]
    inventory = {
        "source_paths": [str(Path(item).expanduser().resolve()) for item in args.source_paths],
        "text_extensions": sorted(SUPPORTED_EXTENSIONS),
        "asset_extensions": sorted(ASSET_EXTENSIONS),
        "sources": [asdict(record) for record in records],
    }
    (project_dir / "source_inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    write_digest(records, project_dir / "source_digest.md")
    extracted = sum(1 for record in records if record.status == "extracted")
    print(f"Ingested {len(records)} sources ({extracted} extracted) into {project_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
